from PyQt5 import QtWidgets, QtCore
from sales.models import Sale, SaleItem
from decimal import Decimal
from datetime import datetime, timedelta
from django.utils import timezone
from django.db.models import Prefetch


class ExportWindow(QtWidgets.QMainWindow):
    def __init__(self):
        super().__init__()

        self.setWindowTitle("Export Sales Report - NASS")
        self.resize(600, 400)
        self.center_window()

        central = QtWidgets.QWidget()
        self.setCentralWidget(central)
        layout = QtWidgets.QVBoxLayout()
        central.setLayout(layout)

        # Title
        title = QtWidgets.QLabel("Export Sales Report")
        title.setAlignment(QtCore.Qt.AlignCenter)
        title.setStyleSheet(
            "font-size: 20px; font-weight: bold; margin: 10px; color: green;"
        )
        layout.addWidget(title)

        # Duration selection group
        duration_group = QtWidgets.QGroupBox("Select Report Duration")
        duration_layout = QtWidgets.QVBoxLayout()
        duration_group.setLayout(duration_layout)

        self.duration_var = QtWidgets.QButtonGroup()

        # Weekly option
        self.weekly_radio = QtWidgets.QRadioButton("Weekly (Last 7 days)")
        self.weekly_radio.setChecked(True)
        self.duration_var.addButton(self.weekly_radio)
        duration_layout.addWidget(self.weekly_radio)

        # Monthly option
        self.monthly_radio = QtWidgets.QRadioButton("Monthly (Last 30 days)")
        self.duration_var.addButton(self.monthly_radio)
        duration_layout.addWidget(self.monthly_radio)

        # Yearly option
        self.yearly_radio = QtWidgets.QRadioButton("Yearly (Last 365 days)")
        self.duration_var.addButton(self.yearly_radio)
        duration_layout.addWidget(self.yearly_radio)

        # Custom date range option
        self.custom_radio = QtWidgets.QRadioButton("Custom Date Range")
        self.duration_var.addButton(self.custom_radio)
        duration_layout.addWidget(self.custom_radio)

        layout.addWidget(duration_group)

        # Custom date range options (initially disabled)
        custom_date_group = QtWidgets.QGroupBox("Custom Date Range (if selected above)")
        custom_date_layout = QtWidgets.QHBoxLayout()
        custom_date_group.setLayout(custom_date_layout)

        custom_date_layout.addWidget(QtWidgets.QLabel("From:"))
        self.from_date = QtWidgets.QDateEdit()
        self.from_date.setDate(QtCore.QDate.currentDate().addDays(-30))
        self.from_date.setCalendarPopup(True)
        self.from_date.setEnabled(False)
        custom_date_layout.addWidget(self.from_date)

        custom_date_layout.addWidget(QtWidgets.QLabel("To:"))
        self.to_date = QtWidgets.QDateEdit()
        self.to_date.setDate(QtCore.QDate.currentDate())
        self.to_date.setCalendarPopup(True)
        self.to_date.setEnabled(False)
        custom_date_layout.addWidget(self.to_date)

        layout.addWidget(custom_date_group)

        # Connect custom radio to enable/disable date fields
        self.custom_radio.toggled.connect(self.toggle_custom_dates)

        # Buttons
        button_layout = QtWidgets.QHBoxLayout()
        layout.addLayout(button_layout)

        export_btn = QtWidgets.QPushButton("Export")
        export_btn.setStyleSheet(
            "background-color: green; color: white; font-weight: bold;"
        )
        export_btn.clicked.connect(self.export_report)
        button_layout.addWidget(export_btn)

        cancel_btn = QtWidgets.QPushButton("Cancel")
        cancel_btn.clicked.connect(self.close)
        button_layout.addWidget(cancel_btn)

        layout.addStretch()

    def center_window(self):
        """Center the window on the screen"""
        screen_geometry = QtWidgets.QApplication.primaryScreen().geometry()
        window_geometry = self.frameGeometry()
        center_point = screen_geometry.center()
        window_geometry.moveCenter(center_point)
        self.move(window_geometry.topLeft())

    def toggle_custom_dates(self):
        """Enable/disable custom date fields based on radio button"""
        is_custom = self.custom_radio.isChecked()
        self.from_date.setEnabled(is_custom)
        self.to_date.setEnabled(is_custom)

    def get_date_range(self):
        """Get the date range based on selected duration"""
        now = timezone.localtime()
        to_dt = now

        if self.weekly_radio.isChecked():
            from_dt = now - timedelta(days=7)
        elif self.monthly_radio.isChecked():
            from_dt = now - timedelta(days=30)
        elif self.yearly_radio.isChecked():
            from_dt = now - timedelta(days=365)
        else:  # Custom
            from_dt = timezone.make_aware(
                datetime.combine(self.from_date.date().toPyDate(), datetime.min.time())
            )
            to_dt = timezone.make_aware(
                datetime.combine(self.to_date.date().toPyDate(), datetime.max.time())
            )

        return from_dt, to_dt

    def export_report(self):
        """Export the sales report based on selected options"""
        file_filter = "Excel Files (*.xlsx)"
        default_name = f"sales_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

        file_path, _ = QtWidgets.QFileDialog.getSaveFileName(
            self, "Save Report As", default_name, file_filter
        )

        if not file_path:
            return

        try:
            from_dt, to_dt = self.get_date_range()

            # Fetch sales data
            reverse_accessor = "items"  # Based on related_name in Sale model
            if not hasattr(Sale, reverse_accessor):
                for f in SaleItem._meta.fields:
                    if (
                        f.get_internal_type() == "ForeignKey"
                        and f.related_model == Sale
                    ):
                        reverse_accessor = f.related_query_name()
                        break

            saleitems_prefetch = Prefetch(
                reverse_accessor,
                queryset=SaleItem.objects.select_related("item"),
                to_attr="prefetched_items",
            )

            sales = (
                Sale.objects.filter(date__range=(from_dt, to_dt))
                .prefetch_related(saleitems_prefetch)
                .order_by("-date")
            )

            self.export_to_excel(file_path, sales, from_dt, to_dt)

            QtWidgets.QMessageBox.information(
                self,
                "Export Successful",
                f"Report exported successfully to:\n{file_path}",
            )
            self.close()

        except ImportError as e:
            if "openpyxl" in str(e):
                QtWidgets.QMessageBox.warning(
                    self,
                    "Missing Dependency",
                    "openpyxl is required for Excel export.\n\n"
                    "Please install it by running:\n"
                    "pip install openpyxl",
                )
            else:
                QtWidgets.QMessageBox.critical(
                    self, "Export Failed", f"An error occurred: {str(e)}"
                )
        except Exception as e:
            QtWidgets.QMessageBox.critical(
                self, "Export Failed", f"An error occurred: {str(e)}"
            )

    def export_to_excel(self, file_path, sales, from_dt, to_dt):
        """Export sales data to Excel file"""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            raise ImportError("openpyxl is required for Excel export")

        wb = Workbook()
        ws = wb.active
        ws.title = "Sales Report"

        # Define styles
        header_fill = PatternFill(
            start_color="90EE90", end_color="90EE90", fill_type="solid"
        )  # Light green
        header_font = Font(bold=True, color="000000", size=12)
        title_font = Font(bold=True, size=14, color="008000")  # Dark green
        summary_fill = PatternFill(
            start_color="E8F5E9", end_color="E8F5E9", fill_type="solid"
        )  # Very light green
        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        # Shop name
        ws.merge_cells("A1:J1")
        shop_cell = ws["A1"]
        shop_cell.value = "REAL GENERAL SUPPLIERS"
        shop_cell.font = Font(bold=True, size=16, color="008000")
        shop_cell.alignment = Alignment(horizontal="center", vertical="center")

        # Title
        ws.merge_cells("A2:J2")
        title_cell = ws["A2"]
        title_cell.value = "NASS - Sales Report"
        title_cell.font = title_font
        title_cell.alignment = Alignment(horizontal="center", vertical="center")

        # Period
        ws.merge_cells("A3:J3")
        period_cell = ws["A3"]
        period_cell.value = (
            f"Period: {timezone.localtime(from_dt).strftime('%Y-%m-%d')} "
            f"to {timezone.localtime(to_dt).strftime('%Y-%m-%d')}"
        )
        period_cell.alignment = Alignment(horizontal="center")

        # Headers
        headers = [
            "Sale ID",
            "Date",
            "Item Name",
            "Size",
            "Quantity",
            "Unit Price (MK)",
            "Item Total (MK)",
            "Cost Price (MK)",
            "Item Profit (MK)",
            "Item Loss (MK)",
        ]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=5, column=col)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

        # Data
        current_row = 6
        total_quantity = Decimal(0)
        total_amount = Decimal(0)
        total_profit = Decimal(0)
        total_loss = Decimal(0)

        for sale in sales:
            sale_items = getattr(sale, "prefetched_items", [])
            if not sale_items:
                sale_items = list(
                    SaleItem.objects.filter(sale=sale).select_related("item")
                )

            for idx, si in enumerate(sale_items):
                item_name = si.item.name if si.item else "(unknown)"
                item_size = si.item.size if si.item and si.item.size else ""
                item_total = (
                    (si.selling_price * si.quantity) if si.selling_price else Decimal(0)
                )

                ws.cell(row=current_row, column=1).value = sale.id if idx == 0 else ""
                ws.cell(row=current_row, column=2).value = (
                    timezone.localtime(sale.date).strftime("%Y-%m-%d %H:%M") if idx == 0 else ""
                )
                ws.cell(row=current_row, column=3).value = item_name
                ws.cell(row=current_row, column=4).value = item_size
                ws.cell(row=current_row, column=5).value = float(si.quantity)
                ws.cell(row=current_row, column=6).value = si.selling_price
                ws.cell(row=current_row, column=7).value = item_total
                ws.cell(row=current_row, column=8).value = si.cost_price
                profit_value = si.profit if si.profit and si.profit >= 0 else Decimal(0)
                loss_value = (-si.profit) if si.profit and si.profit < 0 else Decimal(0)

                ws.cell(row=current_row, column=9).value = profit_value
                ws.cell(row=current_row, column=10).value = loss_value

                # Format numbers
                ws.cell(row=current_row, column=5).number_format = "#,##0.###"
                for col in [6, 7, 8, 9, 10]:
                    ws.cell(row=current_row, column=col).number_format = "#,##0.00"

                # Apply border
                for col in range(1, 11):
                    ws.cell(row=current_row, column=col).border = border

                total_quantity += si.quantity
                total_amount += item_total
                total_profit += profit_value
                total_loss += loss_value
                current_row += 1

        # Summary section
        current_row += 1
        ws.merge_cells(f"A{current_row}:B{current_row}")
        summary_label = ws.cell(row=current_row, column=1)
        summary_label.value = "SUMMARY"
        summary_label.font = Font(bold=True, size=12)
        summary_label.fill = summary_fill

        # Total Items
        current_row += 1
        ws.cell(row=current_row, column=1).value = "Total Items Sold:"
        ws.cell(row=current_row, column=2).value = float(total_quantity)
        ws.cell(row=current_row, column=2).number_format = "#,##0.###"
        ws.cell(row=current_row, column=2).fill = summary_fill

        # Total Amount
        current_row += 1
        ws.cell(row=current_row, column=1).value = "Total Amount:"
        ws.cell(row=current_row, column=2).value = total_amount
        ws.cell(row=current_row, column=2).number_format = "#,##0.00"
        ws.cell(row=current_row, column=2).fill = summary_fill

        # Total Profit
        current_row += 1
        ws.cell(row=current_row, column=1).value = "Total Profit:"
        ws.cell(row=current_row, column=2).value = total_profit
        ws.cell(row=current_row, column=2).number_format = "#,##0.00"
        ws.cell(row=current_row, column=2).fill = summary_fill

        # Total Loss
        current_row += 1
        ws.cell(row=current_row, column=1).value = "Total Loss:"
        ws.cell(row=current_row, column=2).value = total_loss
        ws.cell(row=current_row, column=2).number_format = "#,##0.00"
        ws.cell(row=current_row, column=2).fill = summary_fill

        # Adjust column widths
        ws.column_dimensions["A"].width = 10
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 25
        ws.column_dimensions["D"].width = 12
        ws.column_dimensions["E"].width = 12
        ws.column_dimensions["F"].width = 15
        ws.column_dimensions["G"].width = 15
        ws.column_dimensions["H"].width = 15
        ws.column_dimensions["I"].width = 15
        ws.column_dimensions["J"].width = 15

        wb.save(file_path)
