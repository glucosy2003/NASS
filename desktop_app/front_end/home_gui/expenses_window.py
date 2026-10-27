from decimal import Decimal
from datetime import datetime

from PyQt5 import QtCore, QtWidgets

from django.db.models import Sum
from django.utils import timezone

from front_end.datetime_widgets import local_now_qdatetime
from sales.models import Expense, ServiceEntry


class ExpenseWindow(QtWidgets.QMainWindow):
    """Record shop expenses and review profit/loss (service income - expenses)."""

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Expenses - NASS")
        self.resize(900, 580)
        self.center_window()

        central = QtWidgets.QWidget()
        self.setCentralWidget(central)
        layout = QtWidgets.QVBoxLayout()
        central.setLayout(layout)

        header = QtWidgets.QLabel("Expenses")
        header.setAlignment(QtCore.Qt.AlignCenter)
        header.setStyleSheet(
            "font-size: 20px; font-weight: bold; color: #2c3e50;"
        )
        layout.addWidget(header)

        # Filter row
        filter_layout = QtWidgets.QHBoxLayout()
        layout.addLayout(filter_layout)
        filter_layout.addWidget(QtWidgets.QLabel("From:"))
        self.from_date = QtWidgets.QDateEdit(
            QtCore.QDate.currentDate().addMonths(-1)
        )
        self.from_date.setCalendarPopup(True)
        filter_layout.addWidget(self.from_date)
        filter_layout.addWidget(QtWidgets.QLabel("To:"))
        self.to_date = QtWidgets.QDateEdit(QtCore.QDate.currentDate())
        self.to_date.setCalendarPopup(True)
        filter_layout.addWidget(self.to_date)
        filter_btn = QtWidgets.QPushButton("Filter")
        filter_btn.clicked.connect(self.load_records)
        filter_layout.addWidget(filter_btn)
        export_btn = QtWidgets.QPushButton("Export Report")
        export_btn.setToolTip("Export expense/profit report to spreadsheet")
        export_btn.clicked.connect(self.open_export_window)
        filter_layout.addWidget(export_btn)
        filter_layout.addStretch()

        # Entry form
        form = QtWidgets.QHBoxLayout()
        layout.addLayout(form)
        form.addWidget(QtWidgets.QLabel("Date:"))
        self.date_edit = QtWidgets.QDateTimeEdit(local_now_qdatetime())
        self.date_edit.setCalendarPopup(True)
        self.date_edit.setDisplayFormat("yyyy-MM-dd HH:mm")
        form.addWidget(self.date_edit)

        form.addWidget(QtWidgets.QLabel("Description:"))
        self.description_input = QtWidgets.QLineEdit()
        self.description_input.setPlaceholderText("e.g. Transport")
        form.addWidget(self.description_input, 1)

        form.addWidget(QtWidgets.QLabel("Amount (MK):"))
        self.amount_input = QtWidgets.QLineEdit()
        self.amount_input.setPlaceholderText("0.00")
        self.amount_input.setFixedWidth(100)
        form.addWidget(self.amount_input)

        add_btn = QtWidgets.QPushButton("Add Expense")
        add_btn.setStyleSheet(
            "background-color: #90EE90; color: #2c3e50; font-weight: bold; "
            "padding: 8px 14px; border: none; border-radius: 4px;"
        )
        add_btn.clicked.connect(self.add_expense)
        form.addWidget(add_btn)

        # Table
        self.table = QtWidgets.QTableWidget()
        self.table.setColumnCount(4)
        self.table.setHorizontalHeaderLabels(
            ["Date", "Description", "Amount (MK)", "Delete"]
        )
        self.table.horizontalHeader().setSectionResizeMode(
            0, QtWidgets.QHeaderView.ResizeToContents
        )
        self.table.horizontalHeader().setSectionResizeMode(
            1, QtWidgets.QHeaderView.Stretch
        )
        self.table.horizontalHeader().setSectionResizeMode(
            2, QtWidgets.QHeaderView.ResizeToContents
        )
        self.table.horizontalHeader().setSectionResizeMode(
            3, QtWidgets.QHeaderView.ResizeToContents
        )
        self.table.setEditTriggers(QtWidgets.QAbstractItemView.NoEditTriggers)
        self.table.setSelectionBehavior(QtWidgets.QAbstractItemView.SelectRows)
        self.table.verticalHeader().setDefaultSectionSize(40)
        layout.addWidget(self.table)

        # Summary (profit/loss based on services income)
        self.summary_label = QtWidgets.QLabel()
        self.summary_label.setStyleSheet(
            "font-size: 15px; font-weight: bold; margin-top: 10px; "
            "background-color: #E8F5E9; color: #008000; padding: 10px; "
            "border-radius: 5px; text-align: center;"
        )
        layout.addWidget(self.summary_label)

        self.load_records()

    def center_window(self):
        screen_geometry = QtWidgets.QApplication.primaryScreen().geometry()
        geometry = self.frameGeometry()
        geometry.moveCenter(screen_geometry.center())
        self.move(geometry.topLeft())

    def _period(self):
        from_dt = timezone.make_aware(
            datetime.combine(self.from_date.date().toPyDate(), datetime.min.time())
        )
        to_dt = timezone.make_aware(
            datetime.combine(self.to_date.date().toPyDate(), datetime.max.time())
        )
        return from_dt, to_dt

    def _expense_qs(self):
        from_dt, to_dt = self._period()
        return (
            Expense.objects.filter(deleted=False, date__range=(from_dt, to_dt))
            .order_by("-date")
        )

    def add_expense(self):
        description = self.description_input.text().strip()
        if not description:
            QtWidgets.QMessageBox.warning(
                self, "Error", "Enter a description."
            )
            return
        try:
            amount = Decimal(
                self.amount_input.text().replace("MK", "").replace(",", "").strip()
            )
            if amount <= 0:
                raise ValueError
        except Exception:
            QtWidgets.QMessageBox.warning(
                self, "Error", "Enter a valid amount greater than zero."
            )
            return

        try:
            selected = self.date_edit.dateTime().toPyDateTime()
            date = (
                timezone.make_aware(selected)
                if timezone.is_naive(selected)
                else selected
            )
            Expense.objects.create(
                date=date,
                description=description[:255],
                amount=amount,
            )
        except Exception as exc:  # noqa: BLE001
            QtWidgets.QMessageBox.critical(
                self, "Error", f"Could not save expense: {exc}"
            )
            return

        self.description_input.clear()
        self.amount_input.clear()
        self.load_records()

    def load_records(self):
        self.expenses = list(self._expense_qs())
        self.table.setRowCount(0)
        for row, expense in enumerate(self.expenses):
            self.table.insertRow(row)
            self.table.setItem(
                row,
                0,
                QtWidgets.QTableWidgetItem(
                    timezone.localtime(expense.date).strftime("%Y-%m-%d %H:%M")
                ),
            )
            self.table.setItem(
                row, 1, QtWidgets.QTableWidgetItem(expense.description)
            )
            self.table.setItem(
                row, 2, QtWidgets.QTableWidgetItem(f"MK {expense.amount:,.2f}")
            )

            delete_btn = QtWidgets.QPushButton("Delete")
            delete_btn.clicked.connect(
                lambda _, e=expense: self.delete_expense(e)
            )
            self.table.setCellWidget(row, 3, delete_btn)

        from_dt, to_dt = self._period()
        total_expenses = (
            Expense.objects.filter(
                deleted=False, date__range=(from_dt, to_dt)
            ).aggregate(total=Sum("amount"))["total"]
            or Decimal("0")
        )
        service_income = (
            ServiceEntry.objects.filter(
                deleted=False, date__range=(from_dt, to_dt)
            ).aggregate(total=Sum("total_amount"))["total"]
            or Decimal("0")
        )
        net = service_income - total_expenses
        net_color = "#27ae60" if net >= 0 else "#c0392b"
        self.summary_label.setText(
            f"Expenses: {len(self.expenses)}  ·  Total Expenses: MK {total_expenses:,.2f}"
            f"  ·  Service Income: MK {service_income:,.2f}"
            f"  ·  TOTAL: MK {net:,.2f}"
        )
        self.summary_label.setStyleSheet(
            "font-size: 15px; font-weight: bold; margin-top: 10px; "
            f"background-color: #E8F5E9; color: {net_color}; padding: 10px; "
            "border-radius: 5px; text-align: center;"
        )

    def delete_expense(self, expense):
        answer = QtWidgets.QMessageBox.question(
            self,
            "Confirm delete",
            f"Delete expense '{expense.description}' (MK {expense.amount:,.2f})?",
            QtWidgets.QMessageBox.Yes | QtWidgets.QMessageBox.No,
            QtWidgets.QMessageBox.No,
        )
        if answer == QtWidgets.QMessageBox.Yes:
            expense.deleted = True
            expense.save()
            self.load_records()

    def open_export_window(self):
        file_filter = "Excel Files (*.xlsx)"
        default_name = (
            f"expenses_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        )
        file_path, _ = QtWidgets.QFileDialog.getSaveFileName(
            self, "Save Report As", default_name, file_filter
        )
        if not file_path:
            return
        try:
            from_dt, to_dt = self._period()
            self.export_to_excel(file_path, list(self._expense_qs()), from_dt, to_dt)
            QtWidgets.QMessageBox.information(
                self,
                "Export Successful",
                f"Report exported successfully to:\n{file_path}",
            )
        except ImportError as exc:
            if "openpyxl" in str(exc):
                QtWidgets.QMessageBox.warning(
                    self,
                    "Missing Dependency",
                    "openpyxl is required for Excel export.\n\n"
                    "Please install it by running:\n"
                    "pip install openpyxl",
                )
            else:
                QtWidgets.QMessageBox.critical(
                    self, "Export Failed", f"An error occurred: {str(exc)}"
                )
        except Exception as exc:  # noqa: BLE001
            QtWidgets.QMessageBox.critical(
                self, "Export Failed", f"An error occurred: {str(exc)}"
            )

    def export_to_excel(self, file_path, expenses, from_dt, to_dt):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            raise ImportError("openpyxl is required for Excel export")

        wb = Workbook()
        ws = wb.active
        ws.title = "Expenses Report"

        header_fill = PatternFill(
            start_color="90EE90", end_color="90EE90", fill_type="solid"
        )
        header_font = Font(bold=True, color="000000", size=12)
        title_font = Font(bold=True, size=14, color="008000")
        summary_fill = PatternFill(
            start_color="E8F5E9", end_color="E8F5E9", fill_type="solid"
        )
        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        ws.merge_cells("A1:E1")
        shop_cell = ws["A1"]
        shop_cell.value = "REAL GENERAL SUPPLIERS"
        shop_cell.font = Font(bold=True, size=16, color="008000")
        shop_cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("A2:E2")
        title_cell = ws["A2"]
        title_cell.value = "NASS - Expenses & Profit/Loss Report"
        title_cell.font = title_font
        title_cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("A3:E3")
        period_cell = ws["A3"]
        period_cell.value = (
            f"Period: {timezone.localtime(from_dt).strftime('%Y-%m-%d')} "
            f"to {timezone.localtime(to_dt).strftime('%Y-%m-%d')}"
        )
        period_cell.alignment = Alignment(horizontal="center")

        headers = ["Date", "Description", "Amount (MK)"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=5, column=col)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

        current_row = 6
        for expense in expenses:
            ws.cell(row=current_row, column=1).value = (
                timezone.localtime(expense.date).strftime("%Y-%m-%d %H:%M")
            )
            ws.cell(row=current_row, column=2).value = expense.description
            ws.cell(row=current_row, column=3).value = expense.amount
            ws.cell(row=current_row, column=3).number_format = "#,##0.00"
            for col in range(1, 4):
                ws.cell(row=current_row, column=col).border = border
            current_row += 1

        total_expenses = (
            Expense.objects.filter(
                deleted=False, date__range=(from_dt, to_dt)
            ).aggregate(total=Sum("amount"))["total"]
            or Decimal("0")
        )
        service_income = (
            ServiceEntry.objects.filter(
                deleted=False, date__range=(from_dt, to_dt)
            ).aggregate(total=Sum("total_amount"))["total"]
            or Decimal("0")
        )
        net = service_income - total_expenses

        current_row += 1
        ws.merge_cells(f"A{current_row}:B{current_row}")
        summary_label = ws.cell(row=current_row, column=1)
        summary_label.value = "SUMMARY"
        summary_label.font = Font(bold=True, size=12)
        summary_label.fill = summary_fill

        current_row += 1
        ws.cell(row=current_row, column=1).value = "Expenses count:"
        ws.cell(row=current_row, column=2).value = len(expenses)
        ws.cell(row=current_row, column=2).fill = summary_fill

        current_row += 1
        ws.cell(row=current_row, column=1).value = "Total Expenses:"
        ws.cell(row=current_row, column=2).value = total_expenses
        ws.cell(row=current_row, column=2).number_format = "#,##0.00"
        ws.cell(row=current_row, column=2).fill = summary_fill

        current_row += 1
        ws.cell(row=current_row, column=1).value = "Service Income:"
        ws.cell(row=current_row, column=2).value = service_income
        ws.cell(row=current_row, column=2).number_format = "#,##0.00"
        ws.cell(row=current_row, column=2).fill = summary_fill

        current_row += 1
        ws.cell(row=current_row, column=1).value = "TOTAL (Income - Expenses):"
        ws.cell(row=current_row, column=2).value = net
        ws.cell(row=current_row, column=2).number_format = "#,##0.00"
        ws.cell(row=current_row, column=2).fill = summary_fill

        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 40
        ws.column_dimensions["C"].width = 15

        wb.save(file_path)