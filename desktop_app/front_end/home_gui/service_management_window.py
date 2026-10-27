from decimal import Decimal, InvalidOperation
from datetime import datetime

from PyQt5 import QtCore, QtWidgets

from django.utils import timezone
from django.db.models import Sum

from front_end.datetime_widgets import local_now_qdatetime
from sales.models import Service, ServiceEntry


class ServiceManagementWindow(QtWidgets.QMainWindow):
    """Admin-only window (below Stock Management) with three sections:

    1. Add Service  - manage the service catalog (name + price, no quantity)
    2. Service Entry - sales-style form to record service charges
    3. Services Records - profit/loss-style report with its own summary + export
    """

    def __init__(self, parent=None, admin_user=None, autosync=None):
        super().__init__(parent)
        self.admin_user = admin_user
        self.autosync = autosync
        self.setWindowTitle("Service Management - NASS")
        self.resize(980, 600)
        self.center_window()

        central = QtWidgets.QWidget()
        self.setCentralWidget(central)
        layout = QtWidgets.QVBoxLayout()
        central.setLayout(layout)

        self.tabs = QtWidgets.QTabWidget()
        layout.addWidget(self.tabs)

        self.services_tab = QtWidgets.QWidget()
        self.entry_tab = QtWidgets.QWidget()
        self.records_tab = QtWidgets.QWidget()
        self.tabs.addTab(self.services_tab, "Add Service")
        self.tabs.addTab(self.entry_tab, "Service Entry")
        self.tabs.addTab(self.records_tab, "Services Records")

        self._build_services_tab()
        self._build_entry_tab()
        self._build_records_tab()

        self.load_services()
        self.load_records()

    def center_window(self):
        screen_geometry = QtWidgets.QApplication.primaryScreen().geometry()
        geometry = self.frameGeometry()
        geometry.moveCenter(screen_geometry.center())
        self.move(geometry.topLeft())

    # ====================== 1. ADD SERVICE ======================
    def _build_services_tab(self):
        s_layout = QtWidgets.QVBoxLayout(self.services_tab)

        header = QtWidgets.QLabel("Add Service")
        header.setStyleSheet("font-size: 16px; font-weight: bold;")
        s_layout.addWidget(header)

        hint = QtWidgets.QLabel(
            "Add a service the shop provides, e.g. 'fitting big tyre' at MK 5,000. "
            "Services have name and price only (no quantity)."
        )
        hint.setWordWrap(True)
        hint.setStyleSheet("color: #7f8c8d; font-size: 12px;")
        s_layout.addWidget(hint)

        form = QtWidgets.QHBoxLayout()
        s_layout.addLayout(form)
        form.addWidget(QtWidgets.QLabel("Service name:"))
        self.service_name_input = QtWidgets.QLineEdit()
        self.service_name_input.setPlaceholderText("e.g. fitting big tyre")
        form.addWidget(self.service_name_input, 1)
        form.addWidget(QtWidgets.QLabel("Price (MK):"))
        self.service_price_input = QtWidgets.QLineEdit()
        self.service_price_input.setPlaceholderText("e.g. 5000")
        self.service_price_input.setFixedWidth(110)
        form.addWidget(self.service_price_input)
        add_svc_btn = QtWidgets.QPushButton("Add Service")
        add_svc_btn.setStyleSheet(
            "background-color: #27ae60; color: white; font-weight: bold; "
            "padding: 6px 18px; border: none; border-radius: 4px;"
        )
        add_svc_btn.clicked.connect(self.add_service)
        form.addWidget(add_svc_btn)

        self.services_table = QtWidgets.QTableWidget()
        self.services_table.setColumnCount(4)
        self.services_table.setHorizontalHeaderLabels(
            ["Service", "Price (MK)", "Status", "Actions"]
        )
        header = self.services_table.horizontalHeader()
        header.setSectionResizeMode(0, QtWidgets.QHeaderView.Stretch)
        header.setSectionResizeMode(1, QtWidgets.QHeaderView.ResizeToContents)
        header.setSectionResizeMode(2, QtWidgets.QHeaderView.ResizeToContents)
        header.setSectionResizeMode(3, QtWidgets.QHeaderView.ResizeToContents)
        self.services_table.setEditTriggers(
            QtWidgets.QAbstractItemView.NoEditTriggers
        )
        self.services_table.setSelectionBehavior(
            QtWidgets.QAbstractItemView.SelectRows
        )
        s_layout.addWidget(self.services_table)

    def load_services(self):
        self.services = list(Service.objects.all().order_by("-created_at"))
        self.services_table.setRowCount(0)
        for row, service in enumerate(self.services):
            self.services_table.insertRow(row)
            self.services_table.setItem(
                row, 0, QtWidgets.QTableWidgetItem(service.name)
            )
            self.services_table.setItem(
                row, 1, QtWidgets.QTableWidgetItem(f"MK {service.price:,.2f}")
            )
            self.services_table.setItem(
                row,
                2,
                QtWidgets.QTableWidgetItem(
                    "Deleted" if service.deleted else "Active"
                ),
            )
            btn_box = QtWidgets.QHBoxLayout()
            btn_box.setContentsMargins(0, 0, 0, 0)
            btn_box.setSpacing(4)
            container = QtWidgets.QWidget()
            container.setLayout(btn_box)
            edit_btn = QtWidgets.QPushButton("Edit")
            edit_btn.setFixedHeight(24)
            edit_btn.clicked.connect(
                lambda checked=False, svc=service: self.edit_service(svc)
            )
            btn_box.addWidget(edit_btn)
            del_btn = QtWidgets.QPushButton("Delete")
            del_btn.setFixedHeight(24)
            del_btn.clicked.connect(
                lambda checked=False, svc=service: self.delete_service(svc)
            )
            btn_box.addWidget(del_btn)
            self.services_table.setCellWidget(row, 3, container)

    def _get_price_from_input(self):
        try:
            price = Decimal(self.service_price_input.text().strip())
            if price < 0:
                raise InvalidOperation
            return price
        except (InvalidOperation, ValueError):
            QtWidgets.QMessageBox.warning(
                self, "Error", "Enter a valid price."
            )
            return None

    def add_service(self):
        name = self.service_name_input.text().strip()
        price = self._get_price_from_input()
        if not name:
            QtWidgets.QMessageBox.warning(
                self, "Error", "Enter the service name."
            )
            return
        if price is None:
            return
        if Service.objects.filter(name__iexact=name, deleted=False).exists():
            QtWidgets.QMessageBox.warning(
                self, "Error", "A service with that name already exists."
            )
            return
        Service.objects.create(name=name[:150], price=price)
        self.service_name_input.clear()
        self.service_price_input.clear()
        self.service_name_input.setFocus()
        self.load_services()
        self.reload_entry_services()

    def edit_service(self, service=None):
        if service is None:
            return
        new_name, ok = QtWidgets.QInputDialog.getText(
            self, "Service name", "Name:", text=service.name
        )
        if not ok or not new_name.strip():
            return
        text, ok = QtWidgets.QInputDialog.getText(
            self, "Service price", "Price (MK):",
            text=f"{service.price:.2f}",
        )
        if not ok:
            return
        try:
            price = Decimal(text.strip())
            if price < 0:
                raise InvalidOperation
        except (InvalidOperation, ValueError):
            QtWidgets.QMessageBox.warning(self, "Error", "Enter a valid price.")
            return
        service.name = new_name.strip()[:150]
        service.price = price
        service.deleted = False
        service.save()
        self.load_services()
        self.reload_entry_services()

    def delete_service(self, service=None):
        if service is None:
            return
        answer = QtWidgets.QMessageBox.question(
            self,
            "Confirm delete",
            f"Mark '{service.name}' as deleted? Past entries are kept.",
            QtWidgets.QMessageBox.Yes | QtWidgets.QMessageBox.No,
            QtWidgets.QMessageBox.No,
        )
        if answer == QtWidgets.QMessageBox.Yes:
            service.deleted = True
            service.save()
            self.load_services()
            self.reload_entry_services()

    # ====================== 2. SERVICE ENTRY ======================
    def _build_entry_tab(self):
        e_layout = QtWidgets.QVBoxLayout(self.entry_tab)

        header = QtWidgets.QLabel("Service Entry")
        header.setStyleSheet("font-size: 16px; font-weight: bold;")
        e_layout.addWidget(header)

        hint = QtWidgets.QLabel(
            "Record service charges provided to customers. "
            "Pick a service, add it, then record all charges at once."
        )
        hint.setWordWrap(True)
        hint.setStyleSheet("color: #7f8c8d; font-size: 12px;")
        e_layout.addWidget(hint)

        form = QtWidgets.QHBoxLayout()
        e_layout.addLayout(form)
        form.addWidget(QtWidgets.QLabel("Service:"))
        self.cart_service_combo = QtWidgets.QComboBox()
        self.cart_service_combo.setEditable(True)
        self.cart_service_combo.setInsertPolicy(QtWidgets.QComboBox.NoInsert)
        self.cart_service_combo.setMinimumWidth(320)
        self.cart_service_combo.lineEdit().setPlaceholderText("Search service")
        self.cart_service_combo.currentIndexChanged.connect(
            self._on_cart_service_picked
        )
        form.addWidget(self.cart_service_combo, 1)
        form.addWidget(QtWidgets.QLabel("Price (MK):"))
        self.cart_price_input = QtWidgets.QLineEdit()
        self.cart_price_input.setPlaceholderText("price")
        self.cart_price_input.setFixedWidth(110)
        form.addWidget(self.cart_price_input)

        form.addWidget(QtWidgets.QLabel("Entry date:"))
        self.entry_datetime = QtWidgets.QDateTimeEdit(local_now_qdatetime())
        self.entry_datetime.setCalendarPopup(True)
        self.entry_datetime.setDisplayFormat("yyyy-MM-dd HH:mm")
        form.addWidget(self.entry_datetime)

        add_cart_btn = QtWidgets.QPushButton("Add Service")
        add_cart_btn.setStyleSheet(
            """
            QPushButton {
                background-color: #90EE90;
                color: #2c3e50;
                padding: 6px 16px;
                font-size: 13px;
                font-weight: bold;
                border: none;
                border-radius: 4px;
            }
            QPushButton:hover { background-color: #7CCD7C; }
        """
        )
        add_cart_btn.clicked.connect(self.add_service_to_cart)
        form.addWidget(add_cart_btn)

        self.cart_table = QtWidgets.QTableWidget()
        self.cart_table.setColumnCount(3)
        self.cart_table.setHorizontalHeaderLabels(
            ["Service", "Price (MK)", "Remove"]
        )
        self.cart_table.horizontalHeader().setSectionResizeMode(
            0, QtWidgets.QHeaderView.Stretch
        )
        self.cart_table.horizontalHeader().setSectionResizeMode(
            1, QtWidgets.QHeaderView.ResizeToContents
        )
        self.cart_table.horizontalHeader().setSectionResizeMode(
            2, QtWidgets.QHeaderView.ResizeToContents
        )
        self.cart_table.setEditTriggers(
            QtWidgets.QAbstractItemView.NoEditTriggers
        )
        self.cart_table.setSelectionBehavior(
            QtWidgets.QAbstractItemView.SelectRows
        )
        self.cart_table.verticalHeader().setDefaultSectionSize(40)
        e_layout.addWidget(self.cart_table)

        self.cart_total_label = QtWidgets.QLabel("Total: MK 0.00")
        self.cart_total_label.setStyleSheet(
            "font-size: 18px; font-weight: bold; color: black;"
        )
        e_layout.addWidget(self.cart_total_label)

        record_btn = QtWidgets.QPushButton("Record Service Charge(s)")
        record_btn.setMaximumHeight(32)
        record_btn.setMinimumWidth(300)
        record_btn.setStyleSheet(
            """
            QPushButton {
                background-color: #90EE90;
                color: #2c3e50;
                font-size: 15px;
                font-weight: bold;
                padding: 4px 8px;
                border: none;
                border-radius: 4px;
            }
            QPushButton:hover { background-color: #27ae60; }
            QPushButton:pressed { background-color: #229954; }
        """
        )
        record_btn.clicked.connect(self.record_services)
        e_layout.addWidget(record_btn, 0, QtCore.Qt.AlignCenter)

        e_layout.addStretch()

        self.reload_entry_services()

    def reload_entry_services(self):
        self.cart_services = list(
            Service.objects.filter(deleted=False).order_by("name")
        )
        previous = self.cart_service_combo.currentText()
        self.cart_service_combo.blockSignals(True)
        self.cart_service_combo.clear()
        self.cart_service_combo.addItems(
            [self._display_service(s) for s in self.cart_services]
        )
        index = self.cart_service_combo.findText(previous)
        if index >= 0:
            self.cart_service_combo.setCurrentIndex(index)
        else:
            self.cart_service_combo.setCurrentIndex(-1)
            self.cart_price_input.clear()
        self.cart_service_combo.blockSignals(False)

    @staticmethod
    def _display_service(service):
        return f"{service.name} (MK {service.price:,.2f})"

    def _on_cart_service_picked(self, index):
        if 0 <= index < len(self.cart_services):
            self.cart_price_input.setText(
                f"{self.cart_services[index].price:.2f}"
            )

    def _resolve_cart_service(self):
        text = self.cart_service_combo.currentText().strip()
        for service in self.cart_services:
            if self._display_service(service) == text or service.name == text:
                return service
        return None

    def add_service_to_cart(self):
        service = self._resolve_cart_service()
        if service is None:
            QtWidgets.QMessageBox.warning(
                self, "Error", "Select a service first."
            )
            return
        price_text = self.cart_price_input.text().strip()
        try:
            price = Decimal(price_text)
            if price < 0:
                raise InvalidOperation
        except (InvalidOperation, ValueError):
            QtWidgets.QMessageBox.warning(
                self, "Error", "Enter a valid price."
            )
            return
        row = self.cart_table.rowCount()
        self.cart_table.insertRow(row)
        self.cart_table.setItem(
            row, 0, QtWidgets.QTableWidgetItem(service.name)
        )
        self.cart_table.setItem(
            row, 1, QtWidgets.QTableWidgetItem(f"MK {price:,.2f}")
        )
        remove_btn = QtWidgets.QPushButton("Remove")
        remove_btn.clicked.connect(
            lambda checked=False, r=row: self.remove_cart_row(r)
        )
        self.cart_table.setCellWidget(row, 2, remove_btn)
        self.cart_service_combo.blockSignals(True)
        self.cart_service_combo.setCurrentIndex(-1)
        self.cart_service_combo.setCurrentText("")
        self.cart_service_combo.blockSignals(False)
        self.cart_price_input.clear()
        self.update_cart_total()

    def remove_cart_row(self, row):
        self.cart_table.removeRow(row)
        self.refresh_remove_buttons()
        self.update_cart_total()

    def refresh_remove_buttons(self):
        for r in range(self.cart_table.rowCount()):
            btn = self.cart_table.cellWidget(r, 2)
            if isinstance(btn, QtWidgets.QPushButton):
                btn.clicked.disconnect()
                btn.clicked.connect(
                    lambda checked=False, row=r: self.remove_cart_row(row)
                )

    def update_cart_total(self):
        total = Decimal("0")
        for r in range(self.cart_table.rowCount()):
            item = self.cart_table.item(r, 1)
            if item is None:
                continue
            try:
                text = item.text().replace("MK", "").replace(",", "").strip()
                total += Decimal(text)
            except InvalidOperation:
                continue
        self.cart_total_label.setText(f"Total: MK {total:,.2f}")

    def record_services(self):
        rows = self.cart_table.rowCount()
        if rows == 0:
            QtWidgets.QMessageBox.warning(
                self, "Nothing to record", "Add at least one service first."
            )
            return
        services = []
        for r in range(rows):
            service = None
            name = self.cart_table.item(r, 0).text() if self.cart_table.item(r, 0) else ""
            price_text = (
                self.cart_table.item(r, 1).text().replace("MK", "").replace(",", "").strip()
                if self.cart_table.item(r, 1)
                else ""
            )
            for s in self.cart_services:
                if s.name == name:
                    service = s
                    break
            if service is None:
                QtWidgets.QMessageBox.warning(
                    self, "Error", f"Service '{name}' is no longer available."
                )
                return
            try:
                price = Decimal(price_text)
                if price < 0:
                    raise InvalidOperation
            except (InvalidOperation, ValueError):
                QtWidgets.QMessageBox.warning(
                    self, "Error", f"Invalid price for '{name}'."
                )
                return
            services.append((service, price))

        ServiceEntry.objects.bulk_create(
            [
                ServiceEntry(
                    service=s,
                    date=self._selected_entry_datetime(),
                    quantity=Decimal("1"),
                    unit_price=p,
                    total_amount=p,
                )
                for s, p in services
            ]
        )

        total = sum(p for _, p in services)
        QtWidgets.QMessageBox.information(
            self,
            "Recorded",
            f"{len(services)} service charge(s) recorded (MK {total:,.2f}).",
        )
        self.cart_table.setRowCount(0)
        self.update_cart_total()
        self.tabs.setCurrentWidget(self.records_tab)
        self.load_records()

    def _selected_entry_datetime(self):
        selected = self.entry_datetime.dateTime().toPyDateTime()
        if timezone.is_naive(selected):
            return timezone.make_aware(selected)
        return selected

    # ====================== 3. SERVICES RECORDS ======================
    def _build_records_tab(self):
        r_layout = QtWidgets.QVBoxLayout(self.records_tab)

        header = QtWidgets.QLabel("Services Records")
        header.setStyleSheet("font-size: 16px; font-weight: bold;")
        r_layout.addWidget(header)

        filter_layout = QtWidgets.QHBoxLayout()
        r_layout.addLayout(filter_layout)
        filter_layout.addWidget(QtWidgets.QLabel("From:"))
        self.records_from = QtWidgets.QDateEdit(
            QtCore.QDate.currentDate().addMonths(-1)
        )
        self.records_from.setCalendarPopup(True)
        filter_layout.addWidget(self.records_from)
        filter_layout.addWidget(QtWidgets.QLabel("To:"))
        self.records_to = QtWidgets.QDateEdit(QtCore.QDate.currentDate())
        self.records_to.setCalendarPopup(True)
        filter_layout.addWidget(self.records_to)
        filter_btn = QtWidgets.QPushButton("Filter")
        filter_btn.clicked.connect(self.load_records)
        filter_layout.addWidget(filter_btn)
        export_btn = QtWidgets.QPushButton("Export Report")
        export_btn.setToolTip("Export services report to spreadsheet")
        export_btn.clicked.connect(self.open_export_window)
        filter_layout.addWidget(export_btn)
        erase_all_btn = QtWidgets.QPushButton("Erase All Service Records…")
        erase_all_btn.setToolTip(
            "Soft-delete every service record (requires admin password) and "
            "sync the void to the server. Re-enter services fresh afterwards."
        )
        erase_all_btn.clicked.connect(self.erase_all_records)
        filter_layout.addWidget(erase_all_btn)
        filter_layout.addStretch()

        self.records_table = QtWidgets.QTableWidget()
        self.records_table.setColumnCount(4)
        self.records_table.setHorizontalHeaderLabels(
            ["Date", "Service", "Price (MK)", "Actions"]
        )
        self.records_table.horizontalHeader().setSectionResizeMode(
            0, QtWidgets.QHeaderView.ResizeToContents
        )
        self.records_table.horizontalHeader().setSectionResizeMode(
            1, QtWidgets.QHeaderView.Stretch
        )
        self.records_table.horizontalHeader().setSectionResizeMode(
            2, QtWidgets.QHeaderView.ResizeToContents
        )
        self.records_table.horizontalHeader().setSectionResizeMode(
            3, QtWidgets.QHeaderView.ResizeToContents
        )
        self.records_table.setEditTriggers(
            QtWidgets.QAbstractItemView.NoEditTriggers
        )
        self.records_table.setSelectionBehavior(
            QtWidgets.QAbstractItemView.SelectRows
        )
        self.records_table.verticalHeader().setDefaultSectionSize(40)
        r_layout.addWidget(self.records_table)

        self.records_summary = QtWidgets.QLabel(
            "Services recorded: 0 | Total: MK 0.00"
        )
        self.records_summary.setStyleSheet(
            "font-size: 15px; font-weight: bold; margin-top: 10px; "
            "background-color: #E8F5E9; color: #008000; padding: 10px; "
            "border-radius: 5px;"
        )
        r_layout.addWidget(self.records_summary)

    def _records_queryset(self):
        from_dt = timezone.make_aware(
            datetime.combine(self.records_from.date().toPyDate(), datetime.min.time())
        )
        to_dt = timezone.make_aware(
            datetime.combine(self.records_to.date().toPyDate(), datetime.max.time())
        )
        return (
            ServiceEntry.objects.filter(deleted=False, date__range=(from_dt, to_dt))
            .select_related("service")
            .order_by("-date")
        )

    def load_records(self):
        self.records = list(self._records_queryset())
        self.records_table.setRowCount(0)
        total = Decimal("0")
        for row, entry in enumerate(self.records):
            self.records_table.insertRow(row)
            self.records_table.setItem(
                row,
                0,
                QtWidgets.QTableWidgetItem(
                    timezone.localtime(entry.date).strftime("%Y-%m-%d %H:%M")
                ),
            )
            self.records_table.setItem(
                row, 1, QtWidgets.QTableWidgetItem(str(entry.service))
            )
            self.records_table.setItem(
                row,
                2,
                QtWidgets.QTableWidgetItem(f"MK {entry.total_amount:,.2f}"),
            )
            delete_btn = QtWidgets.QPushButton("Delete")
            delete_btn.setFixedHeight(24)
            delete_btn.clicked.connect(
                lambda checked=False, e=entry, r=row: self.delete_record(e, r)
            )
            self.records_table.setCellWidget(row, 3, delete_btn)
            total += entry.total_amount
        self.records_summary.setText(
            f"Services recorded: {len(self.records)} | Total: MK {total:,.2f}"
        )

    def delete_record(self, entry, row):
        answer = QtWidgets.QMessageBox.question(
            self,
            "Confirm delete",
            f"Delete service record '{entry.service}' (MK {entry.total_amount:,.2f})?",
            QtWidgets.QMessageBox.Yes | QtWidgets.QMessageBox.No,
            QtWidgets.QMessageBox.No,
        )
        if answer != QtWidgets.QMessageBox.Yes:
            return
        try:
            entry = ServiceEntry.objects.filter(pk=entry.pk).first()
            if entry is None or entry.deleted:
                return
            entry.deleted = True
            entry.save()
            self.records_table.removeRow(row)
            self.load_records()
            self._trigger_sync()
            QtWidgets.QMessageBox.information(
                self, "Deleted", "Service record deleted and queued for sync."
            )
        except Exception as exc:  # noqa: BLE001
            QtWidgets.QMessageBox.critical(
                self, "Error", f"Could not delete record: {exc}"
            )

    def erase_all_records(self):
        count = ServiceEntry.objects.filter(deleted=False).count()
        total = (
            ServiceEntry.objects.filter(deleted=False).aggregate(
                total=Sum("total_amount")
            )["total"]
            or 0
        )
        if count == 0:
            QtWidgets.QMessageBox.information(
                self, "Nothing to erase", "There are no active service records."
            )
            return
        answer = QtWidgets.QMessageBox.question(
            self,
            "Erase all service records",
            f"This will delete ALL {count} service records (MK {total:,.2f} in "
            "total) and void them on the server once synced.\n\n"
            "Afterwards you can re-enter new services fresh.\n\n"
            "Continue?",
            QtWidgets.QMessageBox.Yes | QtWidgets.QMessageBox.No,
            QtWidgets.QMessageBox.No,
        )
        if answer != QtWidgets.QMessageBox.Yes:
            return

        from sales.admin_auth import AdminAuthManager

        username = self.admin_user.username if self.admin_user is not None else None
        if not username:
            QtWidgets.QMessageBox.warning(
                self, "Admin required", "An admin account is required to erase records."
            )
            return

        password, ok = QtWidgets.QInputDialog.getText(
            self,
            "Admin password required",
            "Enter your admin password to erase all service records:",
            QtWidgets.QLineEdit.Password,
        )
        if not ok or not password:
            return
        user = AdminAuthManager.authenticate_admin(username, password)
        if user is None:
            QtWidgets.QMessageBox.warning(
                self, "Invalid password", "The password is incorrect."
            )
            return
        if not user.is_admin:
            QtWidgets.QMessageBox.warning(
                self, "Admin required", "Only admin accounts can erase records."
            )
            return

        try:
            updated = ServiceEntry.objects.filter(deleted=False).update(
                deleted=True,
                updated_at=timezone.now(),
            )
            self.load_records()
            self._trigger_sync()
            QtWidgets.QMessageBox.information(
                self,
                "Erased",
                f"{updated} service record(s) erased and queued for sync. "
                "Re-enter new services fresh.",
            )
        except Exception as exc:  # noqa: BLE001
            QtWidgets.QMessageBox.critical(
                self, "Error", f"Could not erase records: {exc}"
            )

    def _trigger_sync(self):
        if getattr(self, "autosync", None) is not None:
            self.autosync.start_sync()

    def open_export_window(self):
        file_filter = "Excel Files (*.xlsx)"
        default_name = (
            f"services_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        )
        file_path, _ = QtWidgets.QFileDialog.getSaveFileName(
            self, "Save Report As", default_name, file_filter
        )
        if not file_path:
            return
        try:
            entries = list(self._records_queryset())
            from_dt = timezone.make_aware(
                datetime.combine(
                    self.records_from.date().toPyDate(), datetime.min.time()
                )
            )
            to_dt = timezone.make_aware(
                datetime.combine(
                    self.records_to.date().toPyDate(), datetime.max.time()
                )
            )
            self.export_to_excel(file_path, entries, from_dt, to_dt)
            QtWidgets.QMessageBox.information(
                self,
                "Export Successful",
                f"Report exported successfully to:\n{file_path}",
            )
        except ImportError as exc:
            if "openpyxl" in str(exc):
                QtWidgets.QMessageBox.warning(
                    self,
                    "Missing Dependency",
                    "openpyxl is required for Excel export.\n\n"
                    "Please install it by running:\n"
                    "pip install openpyxl",
                )
            else:
                QtWidgets.QMessageBox.critical(
                    self, "Export Failed", f"An error occurred: {str(exc)}"
                )
        except Exception as exc:  # noqa: BLE001
            QtWidgets.QMessageBox.critical(
                self, "Export Failed", f"An error occurred: {str(exc)}"
            )

    def export_to_excel(self, file_path, entries, from_dt, to_dt):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            raise ImportError("openpyxl is required for Excel export")

        wb = Workbook()
        ws = wb.active
        ws.title = "Services Report"

        header_fill = PatternFill(
            start_color="90EE90", end_color="90EE90", fill_type="solid"
        )
        header_font = Font(bold=True, color="000000", size=12)
        title_font = Font(bold=True, size=14, color="008000")
        summary_fill = PatternFill(
            start_color="E8F5E9", end_color="E8F5E9", fill_type="solid"
        )
        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        ws.merge_cells("A1:E1")
        shop_cell = ws["A1"]
        shop_cell.value = "REAL GENERAL SUPPLIERS"
        shop_cell.font = Font(bold=True, size=16, color="008000")
        shop_cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("A2:E2")
        title_cell = ws["A2"]
        title_cell.value = "NASS - Services Report"
        title_cell.font = title_font
        title_cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("A3:E3")
        period_cell = ws["A3"]
        period_cell.value = (
            f"Period: {timezone.localtime(from_dt).strftime('%Y-%m-%d')} "
            f"to {timezone.localtime(to_dt).strftime('%Y-%m-%d')}"
        )
        period_cell.alignment = Alignment(horizontal="center")

        headers = ["Date", "Service Name", "Price (MK)", "Notes"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=5, column=col)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

        current_row = 6
        total_amount = Decimal("0")
        for entry in entries:
            ws.cell(row=current_row, column=1).value = (
                timezone.localtime(entry.date).strftime("%Y-%m-%d %H:%M")
            )
            ws.cell(row=current_row, column=2).value = str(entry.service)
            ws.cell(row=current_row, column=3).value = entry.total_amount
            ws.cell(row=current_row, column=4).value = entry.notes or ""
            ws.cell(row=current_row, column=3).number_format = "#,##0.00"
            for col in range(1, 5):
                ws.cell(row=current_row, column=col).border = border
            total_amount += entry.total_amount
            current_row += 1

        current_row += 1
        ws.merge_cells(f"A{current_row}:B{current_row}")
        summary_label = ws.cell(row=current_row, column=1)
        summary_label.value = "SUMMARY"
        summary_label.font = Font(bold=True, size=12)
        summary_label.fill = summary_fill

        current_row += 1
        ws.cell(row=current_row, column=1).value = "Services recorded:"
        ws.cell(row=current_row, column=2).value = len(entries)
        ws.cell(row=current_row, column=2).fill = summary_fill

        current_row += 1
        ws.cell(row=current_row, column=1).value = "Total:"
        ws.cell(row=current_row, column=2).value = total_amount
        ws.cell(row=current_row, column=2).number_format = "#,##0.00"
        ws.cell(row=current_row, column=2).fill = summary_fill

        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["C"].width = 15
        ws.column_dimensions["D"].width = 25
        ws.column_dimensions["E"].width = 30

        wb.save(file_path)